#!/usr/bin/env python3
"""Apply the Stonko manuscript template to a .docx (helper v2, Oct 2026).

Times New Roman 11 pt everywhere, US Letter with 1 in margins, bold 11 pt black headings,
single-spaced title page and Abstract, double-spaced body with a 0.5 in first-line indent and
no space between paragraphs, single-spaced references (6 pt after each) and tables, landscape
table sections preserved, and a prepended title page.

Formatting only: the text of the manuscript is never changed. A style check (em dashes,
P values, confidence intervals, fonts over 11 pt) is printed at the end as warnings.

    python apply_format.py in.docx --title "Full title" --short-title "Running title" \
        --keywords "kw1, kw2" --date "Oct. 2026" \
        --authors "David P. Stonko, MD, MS^1; Courtenay M. Holscher, MD, PhD^2" \
        --affil "Division of Vascular Surgery ... Baltimore, MD, USA" \
        --affil "Department of Surgery, Johns Hopkins Bayview Medical Center, ... Baltimore, MD, USA" \
        --output out.docx
"""
import argparse
import re
import sys

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph

VERSION = "2.0"
FONT = "Times New Roman"
SZ = Pt(11)
BLACK = RGBColor(0, 0, 0)
DEFAULT_AUTHORS = "David P. Stonko, MD, MS^1"
DEFAULT_AFFIL = ("Division of Vascular Surgery and Endovascular Therapy, Department of Surgery, "
                 "Johns Hopkins University School of Medicine, Baltimore, MD, USA")
DEFAULT_FUNDING = "no funding was utilized."
CORRESPONDENCE = ["David P Stonko", "600 N Wolfe St", "Halsted 668", "Baltimore, MD 21287"]
EMAILS = ["David Stonko: dstonko1@jhmi.edu"]

# Heading 1 text (lower case, numbering stripped) -> region of the manuscript
FRONT = {"abstract", "article highlights", "table of contents summary", "structured abstract"}
REFS = {"references", "reference", "literature cited", "bibliography"}
LEGENDS = {"figure legends", "figure legend", "figures", "figure captions"}


def style_run(r, bold=None, italic=None, sup=False):
    """Times New Roman 11 pt on every script (ascii, hAnsi, cs, eastAsia) so theme fonts never win."""
    r.font.name = FONT
    r.font.size = SZ
    rpr = r._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = rpr.makeelement(qn("w:rFonts"), {})
        rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), FONT)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if rf.get(qn(a)) is not None:
            del rf.attrib[qn(a)]
    if bold is not None:
        r.bold = bold
    if italic is not None:
        r.italic = italic
    if sup:
        r.font.superscript = True


def is_heading(p):
    return p.style is not None and p.style.name.startswith("Heading")


def heading_level(p):
    m = re.search(r"(\d+)$", p.style.name)
    return int(m.group(1)) if m else 1


def norm(text):
    return re.sub(r"^[\dIVXivx]+[.)]?\s+", "", text.strip()).strip(" :.").lower()


def all_paragraphs(doc):
    """Body paragraphs in document order, including those python-docx hides (none inside tables)."""
    return [Paragraph(el, doc._body) for el in doc.element.body.iterchildren(qn("w:p"))]


def has_page_break(p):
    return any(br.get(qn("w:type")) == "page" for br in p._element.iter(qn("w:br")))


def has_section_break(p):
    ppr = p._element.pPr
    return ppr is not None and ppr.find(qn("w:sectPr")) is not None


def parse_authors(spec):
    """'Name, MD, MS^1; Name, PhD^2,3' -> [(name, '1'), (name, '2,3')]"""
    out = []
    for part in spec.split(";"):
        part = part.strip()
        if not part:
            continue
        name, _, sup = part.partition("^")
        out.append((name.strip(), sup.strip()))
    return out


def format_body(doc):
    # ---- pages: US Letter, 1 in margins, landscape sections stay landscape
    for s in doc.sections:
        land = s.orientation == WD_ORIENT.LANDSCAPE
        s.page_width, s.page_height = (Inches(11), Inches(8.5)) if land else (Inches(8.5), Inches(11))
        s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Inches(1)

    # ---- styles
    for name in ("Normal", "Body Text", "First Paragraph", "Compact", "Caption", "Table Grid"):
        try:
            st = doc.styles[name]
            st.font.name = FONT
            st.font.size = SZ
        except KeyError:
            pass
    for name in ("Heading 1", "Heading 2", "Heading 3", "Heading 4", "Title", "Caption"):
        try:
            st = doc.styles[name]
            st.font.name = FONT
            st.font.size = SZ
            st.font.color.rgb = BLACK
            st.font.italic = False
            if name != "Caption":
                st.font.bold = True
        except KeyError:
            pass

    paras = all_paragraphs(doc)
    seen_h1 = any(is_heading(p) and heading_level(p) == 1 for p in paras)
    region = "body" if not seen_h1 else None          # unstructured file: treat everything as body
    after_abstract = False
    for p in paras:
        pf = p.paragraph_format
        for r in p.runs:
            style_run(r)
        head = is_heading(p)
        if head:
            for r in p.runs:
                r.font.color.rgb = BLACK
                r.font.bold = True
                r.font.italic = False
            pf.space_before = Pt(0)
            pf.space_after = Pt(0)
            pf.first_line_indent = None
            if heading_level(p) == 1:
                t = norm(p.text)
                if t in FRONT:
                    region = "abstract" if t in ("abstract", "structured abstract") else "front"
                elif t in REFS:
                    region = "refs"
                elif t in LEGENDS:
                    region = "legend"
                else:
                    region = "body"
                # page structure: Abstract on its own page, body starts on a new page,
                # References start on a new page
                if after_abstract or t == "introduction" or region == "refs":
                    pf.page_break_before = True
                after_abstract = region == "abstract"
            pf.line_spacing_rule = (WD_LINE_SPACING.SINGLE if region in ("abstract", "refs")
                                    else WD_LINE_SPACING.DOUBLE)
            continue
        if has_section_break(p) or has_page_break(p):
            continue
        if p.style is not None and p.style.name == "Caption":
            for r in p.runs:
                r.font.color.rgb = BLACK
                r.font.italic = False
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        if region in ("abstract", None):
            pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
        elif region == "refs":
            pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
            pf.space_after = Pt(6)
        elif region in ("front", "legend"):
            pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
        else:                                           # body
            pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
            listy = p._element.pPr is not None and p._element.pPr.numPr is not None
            if p.text.strip() and not listy and p.alignment in (None, 0, 3):
                pf.first_line_indent = Inches(0.5)

    # ---- tables: single spaced, with their caption (paragraph above) and footnote (paragraph below)
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    for r in p.runs:
                        style_run(r)
                    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
                    p.paragraph_format.space_before = Pt(1)
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.first_line_indent = None
        for el in (t._element.getprevious(), t._element.getnext()):
            if el is not None and el.tag == qn("w:p"):
                q = Paragraph(el, doc._body)
                if is_heading(q) or not q.text.strip():
                    continue
                q.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
                q.paragraph_format.space_before = Pt(0)
                q.paragraph_format.space_after = Pt(6)
                q.paragraph_format.first_line_indent = None


def add_title_page(doc, a):
    first = all_paragraphs(doc)[0]

    def line(text="", bold=None, italic=None):
        p = first.insert_paragraph_before("")
        if text:
            style_run(p.add_run(text), bold, italic)
        return p

    p = line()                                           # 1. Prepared for ... date
    style_run(p.add_run("Prepared for: "))
    style_run(p.add_run("***working paper***"), bold=True, italic=True)
    if a.date:
        p.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
        style_run(p.add_run("\t" + a.date))
    line()
    line(a.title, bold=True)                             # title: bold, ALWAYS single spaced
    line("Short Running Title:")
    line(a.short_title)
    line()
    p = line()                                           # authors with superscript affiliations
    authors = parse_authors(a.authors)
    for i, (name, sup) in enumerate(authors):
        style_run(p.add_run(name))
        if sup:
            style_run(p.add_run(sup), sup=True)
        if i < len(authors) - 1:
            style_run(p.add_run("; "))
    line()
    for i, aff in enumerate(a.affil or [DEFAULT_AFFIL], 1):
        line(f"{i} {aff}")
    line()
    line()
    line("Correspondence and reprint requests:", bold=True)
    for t in CORRESPONDENCE:
        line(t)
    line()
    p = line()
    style_run(p.add_run("Source of Funding:"), bold=True)
    style_run(p.add_run(" " + a.funding))
    line()
    line("Emails:", bold=True)
    for t in EMAILS:
        line(t)
    line()
    p = line()
    style_run(p.add_run("Keywords:"), bold=True)
    if a.keywords:
        style_run(p.add_run(" " + a.keywords))
    first.paragraph_format.page_break_before = True      # title page is its own page
    for p in all_paragraphs(doc):                        # single spaced, no paragraph spacing
        if p._element is first._element:
            break
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.first_line_indent = None


def style_check(doc):
    """House-style warnings. Nothing is changed; the author decides."""
    texts = [p.text for p in all_paragraphs(doc)]
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                texts.append(c.text)
    checks = [
        ("em dash (use a comma, parentheses or a new sentence)", r"—"),
        ("P value not in house style (want P=0.03, P=0.003, P<0.001)",
         r"\b[pP]\s+[=<>]\s*\.?\d[\d.]*|\b[pP][=<>]\s+\.?\d[\d.]*|\b[pP]\s*[=<>]\s*\.\d+|\bp\s*[=<>]\s*\d[\d.]*|\b[pP]\s*=\s*0?\.000\b"),
        ("confidence interval not in house style (want 95% CI, 0.39–0.83)",
         r"95%\s*CI[:]?\s*[\[(]?\s*-?\d[\d.]*\s*,\s*-?\d[\d.]*|95%\s*CI\s+-?\d[\d.]*|95%\s*CI,\s*\d[\d.]*\s*-\s*\d[\d.]*"),
    ]
    warned = False
    for label, pat in checks:
        hits = [m for t in texts for m in re.finditer(pat, t)]
        if hits:
            warned = True
            ex = "; ".join(sorted({h.group(0).strip() for h in hits})[:4])
            print(f"  style check: {len(hits)} x {label}. e.g. {ex}")
    big = 0
    for p in all_paragraphs(doc):
        big += sum(1 for r in p.runs if r.font.size is not None and r.font.size > SZ)
    if big:
        warned = True
        print(f"  style check: {big} runs larger than 11 pt")
    if not warned:
        print("  style check: clean")


def main():
    ap = argparse.ArgumentParser(description=f"Stonko manuscript template, helper v{VERSION}")
    ap.add_argument("infile")
    ap.add_argument("--title", required=True)
    ap.add_argument("--short-title", default="")
    ap.add_argument("--keywords", default="")
    ap.add_argument("--date", default="", help='e.g. "Oct. 2026"')
    ap.add_argument("--authors", default=DEFAULT_AUTHORS,
                    help='authors in order, "; " separated, affiliation numbers after ^ '
                         '(e.g. "A. Author, MD^1; B. Author, PhD^2,3")')
    ap.add_argument("--affil", action="append",
                    help="affiliation text without its number; repeat in order of first appearance")
    ap.add_argument("--funding", default=DEFAULT_FUNDING)
    ap.add_argument("--no-title-page", action="store_true", help="restyle only (revising a formatted paper)")
    ap.add_argument("--output", required=True)
    a = ap.parse_args()

    doc = Document(a.infile)
    format_body(doc)
    if not a.no_title_page:
        add_title_page(doc, a)
    doc.save(a.output)
    print("Formatted ->", a.output)
    style_check(doc)


if __name__ == "__main__":
    sys.exit(main())
